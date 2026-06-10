"""
Indeed Canada job listings scraper.

Indeed embeds all job card data as structured JSON in the page source
(window.mosaic.providerData["mosaic-provider-jobcards"]).
We extract that JSON directly — no HTML parsing of individual cards needed.
"""

import argparse
import logging
import os
import re
import json
import time
from typing import Optional

import requests

from config import settings
from exceptions import NoResultsError, ScraperBlockedError, SiteUnavailableError

log = logging.getLogger(__name__)

INDEED_BASE = "https://ca.indeed.com"

# Regex to pull the job cards JSON block out of the page source
_JOBCARDS_PATTERN = re.compile(
    r'window\.mosaic\.providerData\["mosaic-provider-jobcards"\]\s*=\s*(\{.*?\});\s*window\.mosaic\.providerData',
    re.DOTALL,
)


def _build_url(query: str, location: str, start: int) -> str:
    q = query.replace(" ", "+")
    loc = location.replace(" ", "+")
    return f"{settings.indeed_base_url}?q={q}&l={loc}&sort=date&start={start}"


def _session() -> requests.Session:
    """Create a requests session that mimics a browser. Seeds cookies via homepage visit."""
    s = requests.Session()
    s.headers.update({
        "User-Agent": settings.user_agent,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        "Accept-Language": "en-CA,en;q=0.9",
        # Omit Accept-Encoding — let requests use gzip/deflate only (no Brotli dependency).
    })
    try:
        s.get(INDEED_BASE, timeout=settings.request_timeout_s)
    except Exception:
        pass
    return s


def _extract_jobs_from_page(html: str) -> list[dict]:
    """
    Pull the embedded JSON from the page and return the raw job result dicts.
    Returns [] if the JSON block is not found (last page or blocked).
    """
    m = _JOBCARDS_PATTERN.search(html)
    if not m:
        return []
    try:
        data = json.loads(m.group(1))
        return data["metaData"]["mosaicProviderJobCardsModel"]["results"]
    except (json.JSONDecodeError, KeyError):
        return []


def _parse_result(raw: dict) -> dict:
    """Convert one raw job result dict into a clean flat record."""
    salary_snippet = raw.get("salarySnippet") or {}
    salary = salary_snippet.get("text") or None

    job_key = raw.get("jobkey") or ""
    job_url = f"{INDEED_BASE}/viewjob?jk={job_key}" if job_key else None

    return {
        "Job Title":    raw.get("title"),
        "Company":      raw.get("company"),
        "Location":     raw.get("formattedLocation"),
        "Salary":       salary,
        "Posting Date": raw.get("formattedRelativeTime"),
        "Job Key":      job_key or None,
        "Job URL":      job_url,
    }


def _is_blocked(html: str) -> bool:
    lower = html.lower()
    return any(s in lower for s in ["are you a robot", "unusual traffic", "access denied"])


def scrape_jobs(
    query: str = "Mulesoft developer",
    location: str = "Toronto",
    title_filter: Optional[str] = None,
    limit: int = 0,
) -> list[dict]:
    """
    Scrape Indeed Canada for jobs matching query + location.
    title_filter: if set, only keep jobs whose title contains this string (case-insensitive).
    limit=0 means collect all pages.
    Returns list of flat dicts.
    """
    all_jobs: list[dict] = []
    start = 0
    session = _session()

    try:
        while True:
            url = _build_url(query, location, start)
            log.info("Fetching page start=%d: %s", start, url)

            resp = None
            delay = 2
            for attempt in range(1, settings.max_retries + 1):
                try:
                    resp = session.get(url, timeout=settings.request_timeout_s)
                    if resp.status_code == 200:
                        break
                    if resp.status_code in (403, 429, 503):
                        log.warning("HTTP %d on attempt %d/%d — waiting %ds",
                                    resp.status_code, attempt, settings.max_retries, delay)
                        time.sleep(delay)
                        delay *= 2
                        resp = None
                        continue
                    log.warning("Unexpected HTTP %d", resp.status_code)
                    break
                except requests.exceptions.Timeout:
                    log.warning("Timeout on attempt %d/%d", attempt, settings.max_retries)
                    time.sleep(delay)
                    delay *= 2
                except requests.exceptions.ConnectionError as e:
                    raise SiteUnavailableError(f"Cannot reach Indeed: {e}") from e

            if resp is None:
                raise SiteUnavailableError("Indeed did not respond after max retries")

            if _is_blocked(resp.text):
                raise ScraperBlockedError("Indeed returned a bot-detection page")

            raw_results = _extract_jobs_from_page(resp.text)
            if not raw_results:
                log.info("No job data on start=%d — done", start)
                break

            page_jobs = []
            for raw in raw_results:
                record = _parse_result(raw)
                # Apply title filter if requested
                if title_filter:
                    title = record.get("Job Title") or ""
                    if title_filter.lower() not in title.lower():
                        continue
                page_jobs.append(record)
                if limit and (len(all_jobs) + len(page_jobs)) >= limit:
                    break

            all_jobs.extend(page_jobs)
            log.info("Page start=%d | %d matching jobs so far", start, len(all_jobs))

            if limit and len(all_jobs) >= limit:
                break

            if len(raw_results) < settings.page_size:
                break

            start += settings.page_size
            time.sleep(settings.politeness_delay_s)

    except (ScraperBlockedError, SiteUnavailableError):
        raise
    except Exception as e:
        raise SiteUnavailableError(f"Scraper error: {e}") from e

    if not all_jobs:
        raise NoResultsError(
            f"No jobs found matching query='{query}', location='{location}'"
            + (f", title_filter='{title_filter}'" if title_filter else "")
        )

    log.info("Done — %d jobs", len(all_jobs))
    return all_jobs


def save_to_excel(records: list[dict], path: str) -> None:
    """Save job records to a styled Excel file (blue header, frozen row, auto-sized columns)."""
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)

    column_order = [
        "Job Title", "Company", "Location", "Salary",
        "Posting Date", "Job Key", "Job URL",
    ]
    df = pd.DataFrame(records, columns=column_order)
    df.to_excel(path, index=False, sheet_name="Jobs")

    wb = load_workbook(path)
    ws = wb["Jobs"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, cell in enumerate(ws[1], start=1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        max_len = len(str(cell.value or ""))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for c in row:
                max_len = max(max_len, len(str(c.value or "")))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 80)

    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"Saved {len(records)} jobs → {path}")


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s: %(message)s")

    parser = argparse.ArgumentParser(description="Scrape Indeed Canada job listings")
    parser.add_argument("--query",    default="Mulesoft developer", help="Search term")
    parser.add_argument("--location", default="Toronto",            help="City or province")
    parser.add_argument("--filter",   default="mulesoft",           help="Keep only titles containing this word (case-insensitive)")
    parser.add_argument("--no-filter", action="store_true",         help="Disable title filter, return all results")
    parser.add_argument("--limit",    type=int, default=0,          help="Max results (0 = all)")
    parser.add_argument("--output",   default="output/jobs.xlsx",   help="Output Excel path")
    args = parser.parse_args()

    title_filter = None if args.no_filter else args.filter
    print(f"Searching Indeed Canada: query='{args.query}', location='{args.location}'"
          + (f", title must contain '{title_filter}'" if title_filter else ""))

    try:
        jobs = scrape_jobs(args.query, args.location, title_filter, args.limit)
    except NoResultsError as e:
        print(f"No results: {e}")
        raise SystemExit(1)
    except (ScraperBlockedError, SiteUnavailableError) as e:
        print(f"ERROR: {e}")
        raise SystemExit(1)

    save_to_excel(jobs, args.output)

    print(f"\n=== Summary ===")
    print(f"Query:     {args.query}")
    print(f"Location:  {args.location}")
    print(f"Filter:    title contains '{title_filter}'" if title_filter else "Filter:    none")
    print(f"Jobs:      {len(jobs)}")
    print(f"Output:    {args.output}")


if __name__ == "__main__":
    main()
