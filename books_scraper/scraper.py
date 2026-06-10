"""
books.toscrape.com scraper.
Pure server-rendered HTML — uses requests + BeautifulSoup.
Shared by CLI (Excel export) and main.py (FastAPI).
"""

import argparse
import logging
import os
import time
from typing import Optional
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup, Tag

from config import settings
from exceptions import NoResultsError, SiteUnavailableError

log = logging.getLogger(__name__)

RATING_MAP = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


def _session() -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "User-Agent": settings.user_agent,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
    })
    return s


def parse_book(card: Tag) -> dict:
    """Extract title, price, rating, availability from one article.product_pod card."""
    # Title: anchor title attribute holds the full untruncated title
    anchor = card.select_one("h3 > a")
    title = anchor["title"] if anchor else ""

    # Price: strip the £ symbol and convert to float
    price_el = card.select_one("p.price_color")
    price_text = price_el.get_text(strip=True).replace("£", "").replace("Â", "").strip() if price_el else "0"
    try:
        price = float(price_text)
    except ValueError:
        price = 0.0

    # Rating: p.star-rating has a second CSS class that is the word "One"–"Five"
    rating_el = card.select_one("p.star-rating")
    rating_word = rating_el["class"][1] if rating_el and len(rating_el.get("class", [])) > 1 else "One"
    rating = RATING_MAP.get(rating_word, 1)

    # Availability
    avail_el = card.select_one("p.availability")
    availability = avail_el.get_text(strip=True) if avail_el else "Unknown"

    return {
        "Title": title,
        "Price": price,
        "Rating": rating,
        "Availability": availability,
    }


def scrape_books(
    title_filter: str = "",
    min_rating: int = 1,
    max_price: Optional[float] = None,
    availability_filter: str = "all",
    limit: int = 0,
) -> list[dict]:
    """
    Scrape all books from books.toscrape.com and return filtered list.
    limit=0 means no limit (all matching books).
    """
    all_books: list[dict] = []
    url = settings.base_url
    session = _session()
    page = 1

    try:
        while url:
            log.info("Fetching page %d: %s", page, url)

            resp = None
            delay = 2
            for attempt in range(1, settings.max_retries + 1):
                try:
                    resp = session.get(url, timeout=settings.request_timeout_s)
                    if resp.status_code == 200:
                        break
                    log.warning("HTTP %d on attempt %d/%d", resp.status_code, attempt, settings.max_retries)
                    time.sleep(delay)
                    delay *= 2
                    resp = None
                except requests.exceptions.Timeout:
                    log.warning("Timeout on attempt %d/%d", attempt, settings.max_retries)
                    time.sleep(delay)
                    delay *= 2
                except requests.exceptions.ConnectionError as e:
                    raise SiteUnavailableError(f"Cannot reach books.toscrape.com: {e}") from e

            if not resp:
                raise SiteUnavailableError("Site did not respond after max retries")

            soup = BeautifulSoup(resp.text, "lxml")
            cards = soup.select("article.product_pod")

            for card in cards:
                book = parse_book(card)

                # Apply filters
                if title_filter and title_filter.lower() not in book["Title"].lower():
                    continue
                if book["Rating"] < min_rating:
                    continue
                if max_price is not None and book["Price"] > max_price:
                    continue
                if availability_filter == "in_stock" and "stock" not in book["Availability"].lower():
                    continue

                all_books.append(book)

                if limit and len(all_books) >= limit:
                    log.info("Reached limit of %d books", limit)
                    return all_books

            log.info("Page %d | %d matching books so far", page, len(all_books))

            # Next page
            next_el = soup.select_one("li.next > a")
            if next_el:
                # Relative URL — resolve against current page URL
                url = urljoin(url, next_el["href"])
                page += 1
                time.sleep(settings.politeness_delay_s)
            else:
                break

    except SiteUnavailableError:
        raise
    except Exception as e:
        raise SiteUnavailableError(f"Scraper error: {e}") from e

    if not all_books:
        raise NoResultsError("No books matched the given filters")

    log.info("Done — %d books collected", len(all_books))
    return all_books


def save_to_excel(records: list[dict], path: str) -> None:
    """Save book records to a styled Excel file."""
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)

    column_order = ["Title", "Price", "Rating", "Availability"]
    df = pd.DataFrame(records, columns=column_order)
    df.to_excel(path, index=False, sheet_name="Books")

    wb = load_workbook(path)
    ws = wb["Books"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, cell in enumerate(ws[1], start=1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        max_len = len(str(cell.value or ""))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for c in row:
                max_len = max(max_len, len(str(c.value or "")))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 80)

    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"Saved {len(records)} books → {path}")


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s: %(message)s")

    parser = argparse.ArgumentParser(description="Scrape books.toscrape.com")
    parser.add_argument("--title",        default="",               help="Filter by title (substring, case-insensitive)")
    parser.add_argument("--min-rating",   type=int, default=1,      help="Minimum star rating (1–5)")
    parser.add_argument("--max-price",    type=float, default=None, help="Maximum price in £")
    parser.add_argument("--availability", default="all",            choices=["all", "in_stock"], help="Stock filter")
    parser.add_argument("--limit",        type=int, default=0,      help="Max results (0 = all)")
    parser.add_argument("--output",       default="output/books.xlsx")
    args = parser.parse_args()

    try:
        books = scrape_books(args.title, args.min_rating, args.max_price, args.availability, args.limit)
    except NoResultsError as e:
        print(f"No results: {e}")
        raise SystemExit(1)
    except SiteUnavailableError as e:
        print(f"ERROR: {e}")
        raise SystemExit(1)

    save_to_excel(books, args.output)

    print(f"\n=== Summary ===")
    print(f"Title filter: {args.title or '(none)'}")
    print(f"Min rating:   {args.min_rating}")
    print(f"Max price:    £{args.max_price}" if args.max_price else "Max price:    (none)")
    print(f"Books:        {len(books)}")
    print(f"Output:       {args.output}")


if __name__ == "__main__":
    main()
