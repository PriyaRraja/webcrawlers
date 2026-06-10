"""
Florida MQA Dental License Bulk Scraper
========================================
Fetches all dental provider licenses from:
  https://mqa-internet.doh.state.fl.us/MQASearchServices/HealthCareProviders

Default mode  : scrapes results-list pages only (fast — License, Name, Profession, City, Status)
--details flag: also visits each provider's detail page for full info (slow — adds
                Expiration Date, Issue Date, Address, Controlled Substance, Discipline, Complaint)

Usage:
    python scraper.py               # fast mode, all dental professions
    python scraper.py --details     # full detail mode (much slower)
    python scraper.py --profession 701   # single profession only
    python scraper.py --limit 100        # stop after N records (for testing)
"""

import re
import sys
import time
import argparse
import os
import logging
from typing import Optional

import pandas as pd
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

# ── Logging ───────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ── Constants ─────────────────────────────────────────────────────────────────

BASE_URL = "https://mqa-internet.doh.state.fl.us/MQASearchServices/HealthCareProviders"
BASE_ORIGIN = "https://mqa-internet.doh.state.fl.us"
POLITENESS = 0.3   # seconds between detail-page fetches
PAGE_TIMEOUT = 30_000  # ms

DENTAL_PROFESSIONS: dict[str, str] = {
    "701":  "Dental",
    "702":  "Dental Hygienist",
    "704":  "Dental Laboratory",
    "703":  "Dental Radiographer",
    "707":  "Dental Residency Permits",
    "706":  "Dental Teaching Permits",
    "709":  "Dental Temporary Certificate",
    "708":  "Dental-Health Access Dental",
    "9637": "Dental Out-of-State Telehealth",
    "9638": "Dental Teaching Permits Out-of-State Telehealth",
    "717":  "Dental Expert Witness Certificate",
}

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "dental_licenses.xlsx")

# ── Parser helpers ─────────────────────────────────────────────────────────────

_CLICK_NOISE = re.compile(r'\s*-?\s*Click on.*$', re.IGNORECASE)
_HELP_TEXT   = re.compile(r'if further information is needed', re.IGNORECASE)


def _clean(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    value = _CLICK_NOISE.sub("", value).strip().rstrip("/").strip()
    if _HELP_TEXT.search(value):
        return None
    return value or None


def _dt_dd_map(soup: BeautifulSoup) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for dt in soup.find_all("dt"):
        label = dt.get_text(" ", strip=True)
        if not label:
            continue
        dd = dt.find_next_sibling("dd")
        if dd:
            mapping[label] = dd.get_text(" ", strip=True)
    return mapping


def _match(mapping: dict[str, str], *prefixes: str) -> Optional[str]:
    for key, val in mapping.items():
        for prefix in prefixes:
            if key.strip().lower().startswith(prefix.lower()):
                return val
    return None


def parse_detail_page(html: str) -> dict:
    """Extract all fields from a provider detail page."""
    soup = BeautifulSoup(html, "html.parser")
    m = _dt_dd_map(soup)

    name: Optional[str] = None
    for h3 in soup.find_all("h3"):
        text = h3.get_text(" ", strip=True)
        if not text.lower().startswith("license number"):
            name = text
            break

    license_number: Optional[str] = None
    for h3 in soup.find_all("h3"):
        text = h3.get_text(" ", strip=True)
        if text.lower().startswith("license number"):
            parts = text.split(":", 1)
            if len(parts) == 2:
                license_number = parts[1].strip()
            break
    if not license_number:
        license_number = _match(m, "License")

    return {
        "Name":                          _clean(name),
        "License Number":                _clean(license_number),
        "Profession":                    _clean(_match(m, "Profession")),
        "License Status":                _clean(_match(m, "License Status")),
        "Expiration Date":               _clean(_match(m, "License Expiration Date", "Expiration Date")),
        "Original Issue Date":           _clean(_match(m, "License Original Issue Date", "Original Issue Date")),
        "City":                          None,
        "Address":                       _clean(_match(m, "Address of Record", "Primary Address", "Address")),
        "Controlled Substance":          _clean(_match(m, "Controlled Substance")),
        "Discipline on File":            _clean(_match(m, "Discipline on File")),
        "Public Complaint":              _clean(_match(m, "Public Complaint")),
        "Detail URL":                    None,
    }


def parse_results_row(row_el, profession_label: str) -> tuple[dict, Optional[str]]:
    """
    Extract the 5 columns from a results-table row.
    Returns (record_dict, detail_url_or_None).
    """
    cells = row_el.find_all("td")
    if len(cells) < 5:
        return {}, None

    link = cells[0].find("a")
    detail_url = (BASE_ORIGIN + link["href"]) if link and link.get("href") else None

    record = {
        "License Number": cells[0].get_text(strip=True),
        "Name":           cells[1].get_text(strip=True),
        "Profession":     cells[2].get_text(strip=True) or profession_label,
        "City":           cells[3].get_text(strip=True),
        "License Status": cells[4].get_text(strip=True),
        "Expiration Date":      None,
        "Original Issue Date":  None,
        "Address":              None,
        "Controlled Substance": None,
        "Discipline on File":   None,
        "Public Complaint":     None,
        "Detail URL":           detail_url,
    }
    return record, detail_url

# ── Scraper core ───────────────────────────────────────────────────────────────

def search_profession(page, profession_value: str):
    """Navigate to search page and submit for the given profession code."""
    page.goto(BASE_URL)
    page.wait_for_load_state("networkidle")

    body = page.inner_text("body").lower()
    if "maintenance" in body and not page.query_selector("form"):
        raise RuntimeError("MQA site appears to be under maintenance")

    page.select_option("#ProfessionDD", value=profession_value)
    page.click('input[type="submit"][value="Search"]')
    page.wait_for_load_state("networkidle")


def collect_rows(page) -> tuple[list, int]:
    """
    Parse the current results page.
    Returns (list of BeautifulSoup <tr> elements, total_count).
    """
    soup = BeautifulSoup(page.content(), "html.parser")
    table = soup.find("table")
    rows = table.find_all("tr")[1:] if table else []  # skip header row

    total = 0
    for txt in soup.find_all(string=lambda t: t and "Search Results Total:" in t):
        m = re.search(r'(\d[\d,]*)', txt)
        if m:
            total = int(m.group(1).replace(",", ""))
    return rows, total


def go_next_page(page) -> bool:
    """Click the next-page link if present. Returns True if navigated."""
    soup = BeautifulSoup(page.content(), "html.parser")
    next_link = soup.find("a", string=lambda t: t and t.strip() in ("»", "Next", ">"))
    if not next_link or not next_link.get("href"):
        return False
    url = BASE_ORIGIN + next_link["href"] if next_link["href"].startswith("/") else next_link["href"]
    try:
        page.goto(url)
        page.wait_for_load_state("networkidle")
        return True
    except PWTimeout:
        log.warning("Timeout navigating to next page: %s", url)
        return False


def scrape_with_details(page, profession_value: str, profession_label: str, limit: int) -> list[dict]:
    """Scrape results + visit each detail page. Slow but comprehensive."""
    records = []
    page_num = 0

    try:
        search_profession(page, profession_value)
    except Exception as e:
        log.error("Search failed for %s: %s", profession_label, e)
        return []

    while True:
        page_num += 1
        rows, total = collect_rows(page)
        if page_num == 1:
            log.info("  [%s] %d total records, fetching detail pages...", profession_label, total)

        for row in rows:
            if limit and len(records) >= limit:
                return records

            _, detail_url = parse_results_row(BeautifulSoup(str(row), "html.parser"), profession_label)
            if not detail_url:
                continue

            try:
                page.goto(detail_url)
                page.wait_for_load_state("networkidle")
                record = parse_detail_page(page.content())
                record["Detail URL"] = detail_url
                records.append(record)
                page.go_back()
                page.wait_for_load_state("networkidle")
                time.sleep(POLITENESS)
            except PWTimeout:
                log.warning("  Timeout on detail page: %s — skipping", detail_url)
                try:
                    page.go_back()
                    page.wait_for_load_state("networkidle")
                except Exception:
                    search_profession(page, profession_value)
                continue
            except Exception as e:
                log.warning("  Error on detail page: %s — %s", detail_url, e)
                continue

            if len(records) % 50 == 0:
                log.info("  [%s] Page %d | %d records collected", profession_label, page_num, len(records))

        if not go_next_page(page):
            break

    return records


def scrape_list_only(page, profession_value: str, profession_label: str, limit: int) -> list[dict]:
    """Scrape only the results list pages. Fast."""
    records = []
    page_num = 0

    try:
        search_profession(page, profession_value)
    except Exception as e:
        log.error("Search failed for %s: %s", profession_label, e)
        return []

    while True:
        page_num += 1
        rows, total = collect_rows(page)
        if page_num == 1:
            log.info("  [%s] %d total records across ~%d pages", profession_label, total, (total // 20) + 1)

        soup_rows = [BeautifulSoup(str(r), "html.parser") for r in rows]
        for soup_row in soup_rows:
            if limit and len(records) >= limit:
                return records
            record, _ = parse_results_row(soup_row, profession_label)
            if record:
                records.append(record)

        if len(records) % 500 == 0 and records:
            log.info("  [%s] Page %d | %d records", profession_label, page_num, len(records))

        if not go_next_page(page):
            break

    log.info("  [%s] Done — %d records", profession_label, len(records))
    return records

# ── Excel export ───────────────────────────────────────────────────────────────

COLUMN_ORDER = [
    "License Number", "Name", "Profession", "License Status",
    "City", "Expiration Date", "Original Issue Date",
    "Address", "Controlled Substance", "Discipline on File",
    "Public Complaint", "Detail URL",
]


def save_to_excel(records: list[dict], path: str):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    df = pd.DataFrame(records)

    # Ensure all columns present, in order
    for col in COLUMN_ORDER:
        if col not in df.columns:
            df[col] = None
    df = df[COLUMN_ORDER]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Dental Licenses")
        ws = writer.sheets["Dental Licenses"]

        # Bold header
        from openpyxl.styles import Font, PatternFill, Alignment
        header_font   = Font(bold=True, color="FFFFFF")
        header_fill   = PatternFill("solid", fgColor="2E75B6")
        header_align  = Alignment(horizontal="center", vertical="center")

        for cell in ws[1]:
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = header_align

        ws.row_dimensions[1].height = 20

        # Auto-size columns
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 60)

        # Freeze header row
        ws.freeze_panes = "A2"

    log.info("Saved %d records → %s", len(records), path)

# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="FL MQA Dental License Scraper")
    parser.add_argument("--details",    action="store_true",
                        help="Visit each provider's detail page (much slower, adds expiry/address/discipline)")
    parser.add_argument("--profession", metavar="CODE",
                        help="Scrape only this profession code (e.g. 701 for Dental)")
    parser.add_argument("--limit",     type=int, default=0,
                        help="Stop after N records per profession (0 = no limit; useful for testing)")
    parser.add_argument("--output",    default=OUTPUT_FILE,
                        help=f"Output Excel file path (default: {OUTPUT_FILE})")
    args = parser.parse_args()

    professions = (
        {args.profession: DENTAL_PROFESSIONS.get(args.profession, args.profession)}
        if args.profession
        else DENTAL_PROFESSIONS
    )

    mode = "DETAIL PAGES" if args.details else "RESULTS LIST ONLY"
    log.info("Mode: %s | Professions: %d | Limit: %s",
             mode, len(professions), args.limit or "none")

    all_records: list[dict] = []

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context()
        page = context.new_page()
        page.set_default_timeout(PAGE_TIMEOUT)

        for code, label in professions.items():
            log.info("Searching profession: %s (%s)", label, code)
            try:
                if args.details:
                    records = scrape_with_details(page, code, label, args.limit)
                else:
                    records = scrape_list_only(page, code, label, args.limit)
                all_records.extend(records)
            except Exception as e:
                log.error("Failed profession %s (%s): %s", label, code, e)
                continue

        browser.close()

    if not all_records:
        log.error("No records collected — nothing to save")
        sys.exit(1)

    save_to_excel(all_records, args.output)
    log.info("Done. Total records: %d", len(all_records))


if __name__ == "__main__":
    main()
