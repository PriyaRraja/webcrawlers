"""
Redfin Property Listings Bulk Scraper
======================================
Fetches property listings from Redfin for a given city/state
and exports all fields to a styled Excel file.

Uses Redfin's internal Stingray JSON API — no HTML parsing, no browser.

Usage:
    python scraper.py --city Miami --state FL
    python scraper.py --city Orlando --state FL --type house
    python scraper.py --city Tampa --state FL --status sold --min-beds 3
    python scraper.py --city Miami --state FL --max-price 800000
    python scraper.py --city Miami --state FL --limit 50   # for testing
    python scraper.py --city Miami --state FL --output my_results.xlsx
"""

import sys
import os
import argparse
import logging

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

from redfin_client import search_city
from models import PROPERTY_TYPE_CODES, STATUS_CODES
from exceptions import CityNotFoundError, RateLimitedError, SiteUnavailableError

# ── Logging ───────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ── Excel column order ────────────────────────────────────────────────────────

COLUMN_ORDER = [
    "Address", "Unit", "City", "State", "Zip", "Neighborhood",
    "Price", "Beds", "Baths", "Full Baths", "Sqft", "Price/Sqft",
    "Lot Size (sqft)", "Year Built", "Stories",
    "Property Type", "MLS Status", "Days on Market",
    "HOA Fee", "Garage Spaces", "Parking Spaces", "Has Pool",
    "Latitude", "Longitude",
    "MLS Number", "Property ID", "Listing ID",
    "Listing Agent", "Listing Broker",
    "Key Features", "Description",
    "Has Virtual Tour", "Has 3D Tour", "Is New Construction", "Is Hot Listing",
    "Redfin URL",
]

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")


def save_to_excel(records: list[dict], path: str):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    df = pd.DataFrame(records)

    # Ensure all expected columns exist in order
    for col in COLUMN_ORDER:
        if col not in df.columns:
            df[col] = None
    df = df[COLUMN_ORDER]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Properties")
        ws = writer.sheets["Properties"]

        # Styled header row
        header_font  = Font(bold=True, color="FFFFFF")
        header_fill  = PatternFill("solid", fgColor="1A6496")
        header_align = Alignment(horizontal="center", vertical="center")

        for cell in ws[1]:
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = header_align

        ws.row_dimensions[1].height = 22

        # Auto-size columns
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 60)

        # Freeze header
        ws.freeze_panes = "A2"

    log.info("Saved %d records → %s", len(records), path)


def main():
    parser = argparse.ArgumentParser(description="Redfin Property Listings Scraper")
    parser.add_argument("--city",      required=True, help="City name (e.g. Miami)")
    parser.add_argument("--state",     required=True, help="2-letter state code (e.g. FL)")
    parser.add_argument("--type",      default="all",
                        choices=list(PROPERTY_TYPE_CODES.keys()),
                        help="Property type filter (default: all)")
    parser.add_argument("--status",    default="for_sale",
                        choices=list(STATUS_CODES.keys()),
                        help="Listing status (default: for_sale)")
    parser.add_argument("--min-beds",  type=int,   default=None, help="Minimum bedrooms")
    parser.add_argument("--min-baths", type=float, default=None, help="Minimum bathrooms")
    parser.add_argument("--min-price", type=int,   default=None, help="Minimum price")
    parser.add_argument("--max-price", type=int,   default=None, help="Maximum price")
    parser.add_argument("--min-year-built", type=int, default=None, help="Minimum year built (e.g. 2021)")
    parser.add_argument("--limit",     type=int,   default=0,
                        help="Stop after N records (0 = no limit; use for testing)")
    parser.add_argument("--output",    default=None,
                        help="Output Excel path (default: output/{city}_{state}.xlsx)")
    args = parser.parse_args()

    city    = args.city.strip()
    state   = args.state.strip().upper()
    output  = args.output or os.path.join(OUTPUT_DIR, f"{city.replace(' ', '_')}_{state}.xlsx")

    prop_code   = PROPERTY_TYPE_CODES[args.type]
    status_code = STATUS_CODES[args.status]

    log.info(
        "Searching: %s, %s | Type: %s | Status: %s | Limit: %s",
        city, state, args.type, args.status, args.limit or "none",
    )

    try:
        records, total = search_city(
            city=city,
            state=state,
            property_type_code=prop_code,
            status_code=status_code,
            min_beds=args.min_beds,
            min_baths=args.min_baths,
            min_price=args.min_price,
            max_price=args.max_price,
            min_year_built=args.min_year_built,
            limit=args.limit,
        )
    except CityNotFoundError as e:
        log.error("City not found: %s", e)
        sys.exit(1)
    except RateLimitedError as e:
        log.error("Rate limited by Redfin: %s", e)
        sys.exit(1)
    except SiteUnavailableError as e:
        log.error("Redfin unavailable: %s", e)
        sys.exit(1)

    if not records:
        log.error("No listings found — nothing to save")
        sys.exit(1)

    save_to_excel(records, output)
    log.info("Done. Total: %d records", len(records))


if __name__ == "__main__":
    main()
